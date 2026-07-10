"""
agents/document_generator.py

The Document Generator Agent: takes the ReviewedDocument produced by
ReviewerAgent and writes it to a .docx file on disk.

Call tree (same hybrid-architecture discipline as the other agents):

    DocumentGeneratorAgent
    │
    ├── generate_docx()              
    │       ├── build_document()    
    │       │       ├── add_title()          
    │       │       └── add_section()        
    │       └── save_document()      
    │
    └── get_output_path()            

"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from agents.reviewer import ReviewedDocument
from utils.file_utils import generate_safe_filename, resolve_output_path, get_relative_output_path

if TYPE_CHECKING:
    pass

logger = logging.getLogger(__name__)


class DocumentGeneratorAgentError(Exception):
    """Raised when the Document Generator cannot write the .docx file."""


class DocumentGeneratorAgent:

    # PUBLIC ENTRYPOINT
    def generate_docx(self, reviewed_document: ReviewedDocument) -> str:
        output_path = self.get_output_path(reviewed_document.document_type)
        doc = self.build_document(reviewed_document)
        self.save_document(doc, output_path)

        relative_path = get_relative_output_path(output_path)
        logger.info(
            "Document saved: type=%s, path=%s, sections=%d",
            reviewed_document.document_type,
            relative_path,
            len(reviewed_document.sections),
        )
        return relative_path


    # get_output_path — PURE PYTHON

    def get_output_path(self, document_type: str) -> Path:

        filename = generate_safe_filename(document_type)
        return resolve_output_path(filename)

    # build_document — PURE PYTHON

    def build_document(self, reviewed_document: ReviewedDocument) -> Document:
        """

        Constructs a python-docx Document object from the reviewed
        sections: a title page heading followed by each section's
        heading + body paragraph.
        """
        doc = Document()

        self._add_title(doc, reviewed_document)

        for section in reviewed_document.sections:
            self._add_section(doc, section.section_name, section.content)

        return doc

    # save_document — PURE PYTHON
    def save_document(self, doc: Document, output_path: Path) -> None:
        """
        Writes the python-docx Document to disk at the resolved path.
        """
        try:
            doc.save(str(output_path))
        except OSError as exc:
            logger.error("Failed to save document to %s: %s", output_path, exc)
            raise DocumentGeneratorAgentError(
                f"Could not write .docx file to '{output_path}': {exc}"
            ) from exc

   
    # PRIVATE HELPERS — all PURE PYTHON
   
    def _add_title(self, doc: Document, reviewed_document: ReviewedDocument) -> None:
        """
      

        Adds a centered, styled title heading as the first element of
        the document, built from the document type and the original request.
        """
        # Format: "Business Proposal" from "business_proposal"
        display_type = reviewed_document.document_type.replace("_", " ").title()

        title_para = doc.add_heading(display_type, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Add a subtitle paragraph with a condensed version of the request
        subtitle_text = reviewed_document.original_request
        if len(subtitle_text) > 120:
            subtitle_text = subtitle_text[:117].rsplit(" ", 1)[0] + "..."

        subtitle = doc.add_paragraph(subtitle_text)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle.runs[0] if subtitle.runs else subtitle.add_run(subtitle_text)
        subtitle_run.italic = True
        subtitle_run.font.size = Pt(11)
        subtitle_run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

        doc.add_paragraph()  # blank spacer before sections

    def _add_section(self, doc: Document, section_name: str, content: str) -> None:
        """

        Appends one section (heading + body) to the document.
        """
        doc.add_heading(section_name, level=1)

        # Split content on double newlines to produce separate paragraphs.
        paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]
        if not paragraphs:
            paragraphs = [content.strip()]

        for paragraph_text in paragraphs:
            para = doc.add_paragraph()
            # Preserve single newlines within a paragraph as line breaks.
            lines = paragraph_text.split("\n")
            for i, line in enumerate(lines):
                run = para.add_run(line)
                run.font.size = Pt(11)
                if i < len(lines) - 1:
                    run.add_break()

        doc.add_paragraph()  # blank spacer between sections
